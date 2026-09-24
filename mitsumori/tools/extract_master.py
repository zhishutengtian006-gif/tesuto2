#!/usr/bin/env python3
"""見積Excel(.xlsm) → Web見積システム用 単価マスタ(js/master-data.js) 抽出ツール

使い方:
    pip install openpyxl
    python3 tools/extract_master.py 見積書.xlsm            # js/master-data.js を上書き
    python3 tools/extract_master.py 見積書.xlsm -o out.js

Excel側で単価・掛率・商品を改定したら、このスクリプトを再実行するだけで
Web側の単価マスタが更新されます（数量式・画面はそのまま）。

抽出対象シート:
    契約内訳  … 工事明細（標準/OP/付帯/諸経費/設計）の単価・掛率・端数桁
    ①大工    … 大工手間単価
    ②ｻｯｼ    … サッシ・玄関ドア単価表と標準セット
    ③既建    … 既製建具単価表・デザイン加算・表示錠と標準セット
    ④造建    … 造作建具単価表
    ⑤固棚    … 固定棚単価表と標準数
    ⑥可棚    … 可動棚単価表（1ブロック目）
    ⑦電気    … 電気の標準数量と単価
    ⑧配管    … 給排水の標準数量と単価
"""
import argparse
import datetime
import json
import os
import re
import sys
import warnings

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl が必要です: pip install openpyxl")

warnings.filterwarnings("ignore")


def num(x):
    """数値なら float/int、それ以外は None"""
    if isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        return int(x) if float(x).is_integer() else float(x)
    return None


def txt(x):
    if x is None:
        return ""
    s = str(x).replace("\n", " ").strip()
    return re.sub(r"\s+", " ", s.replace("　", " ")).strip()


def is_formula(x):
    return isinstance(x, str) and x.startswith("=")


class Book:
    def __init__(self, path):
        self.f = openpyxl.load_workbook(path, keep_vba=False)
        self.v = openpyxl.load_workbook(path, data_only=True, keep_vba=False)

    def cell(self, sheet, ref):
        """(式, 計算済み値)"""
        return self.f[sheet][ref].value, self.v[sheet][ref].value

    def val(self, sheet, ref):
        return self.v[sheet][ref].value

    def fml(self, sheet, ref):
        v = self.f[sheet][ref].value
        if hasattr(v, "text"):  # ArrayFormula
            return v.text
        return v


# ---------------------------------------------------------------- 契約内訳
SEC_TAGS = ("基", "OP", "付帯", "諸", "設計")
# 明細ではなく数量計算用の補助行（面積の中間値など）
HELPER_ROWS = {430, 436, 441, 481, 527, 535, 536, 537, 539, 606, 749}


def extract_kouji(b):
    S = "契約内訳"
    ws = b.f[S]
    rows = []
    sec_no, sec_name, group = "", "", ""
    pending = []
    for r in range(105, 1054):
        g = lambda c: b.fml(S, f"{c}{r}")
        gv = lambda c: b.val(S, f"{c}{r}")
        A, D, E, F, H = gv("A"), gv("D"), txt(gv("E")), txt(gv("F")), txt(gv("H"))
        # 小計行 → 直前ブロックに区分タグを付与
        if A in SEC_TAGS and ("小計" in E):
            for it in pending:
                it["tag"] = A
            pending = []
            continue
        if D not in (None, "") and E and not H:
            sec_no, group = txt(D), ""
            sec_name = f"{E}（{F}）" if (E == "木工事" and F) else E
            continue
        if r in HELPER_ROWS:
            continue
        Sf, Sv = g("S"), gv("S")
        Lf, Lv = g("L"), gv("L")
        price = num(Sv)
        has_price = price is not None or is_formula(Sf)
        if not H and not has_price:
            if E or F:
                group = E or F
            continue
        if not (E or F):
            # 保証費など名称がP列(備考)にだけある原価行
            if txt(gv("P")) and price is not None:
                E = txt(gv("P"))
            else:
                continue
        Tf, Tv = g("T"), gv("T")
        Z = num(gv("Z"))
        ab = g("AB") or ""
        ab_round = -3 if (isinstance(ab, str) and ",-3)" in ab) else -2
        it = {
            "row": r,
            "sec": f"{sec_no} {sec_name}".strip(),
            "group": group,
            "code": txt(gv("C")),
            "name": E,
            "spec": F,
            "h": H,  # 表示単位（"式"なら一式計上）
            "unit": txt(gv("M")),
            "price": price,
            "net": num(Tv),
            "rate2": num(gv("V")),
            "z": int(Z) if Z is not None else 1,
            "abRound": ab_round,
            "note": txt(gv("P")) if not is_formula(g("P")) else "",
            "pnum": num(gv("P")),  # P列の補助数値（標準台数・既定数量など）
            "tag": None,
            # 参考（テスト用）: サンプルExcelの計算値
            "xl": {"qty": num(Lv), "sell": num(gv("J")), "cost": num(gv("O"))},
        }
        if is_formula(Sf):
            it["priceFormula"] = Sf
        if is_formula(Lf):
            it["qtyFormula"] = Lf
        elif num(Lf) is not None:
            it["qtyConst"] = num(Lf)
        rows.append(it)
        pending.append(it)
    return rows


def extract_kouji_extra(b):
    S = "契約内訳"
    siding = {}
    for r in range(516, 535):
        a, h = num(b.val(S, f"AG{r}")), num(b.val(S, f"AH{r}"))
        if a is not None or h is not None:
            siding[r] = {"normal": a, "pg": h}
    elec_base = {}
    for c in range(33, 37):
        k = b.v[S].cell(row=600, column=c).value
        vv = b.v[S].cell(row=601, column=c).value
        if k:
            elec_base[txt(k)] = num(vv)
    def alt(ref):
        """=IF(条件,A,B) の A,B を取り出す（PG用単価と通常単価）"""
        m = re.search(r',(\d+),(\d+)\)\s*$', str(b.fml(S, ref) or ""))
        return [int(m.group(1)), int(m.group(2))] if m else None
    return {
        "precutPG": alt("S154"),        # [PG平屋, PG2階建]
        "zatsuKanamono": alt("S577"),   # [PG, その他]
        "siding": siding,
        "elecBase": elec_base,
        "kasetsuDenki": num(b.val(S, "AG968")),
        "kansenHikikomi": num(b.val(S, "AG987")),
        "okugaiSwitch": num(b.val(S, "AG988")),
        "ippanKanriPerTsubo": num(b.val(S, "P1022")),
        "genbaKeihiPerTsubo": num(b.val(S, "P1023")),
        "jibanChosa": num(b.val(S, "P947")),
        "jibanTatekaeRate": num(b.val(S, "R947")),
        "design": {
            "kouzou": [num(b.val(S, "U1039")), num(b.val(S, "W1039"))],
            "gaihi": [num(b.val(S, "U1040")), num(b.val(S, "W1040"))],
            "belsDoc": num(b.val(S, "U1041")),
            "chokiDoc": num(b.val(S, "U1042")),
            "kakunin": [num(b.val(S, "U1043")), num(b.val(S, "W1043"))],
            "seinou": num(b.val(S, "U1044")),
            "bels": [num(b.val(S, "U1045")), num(b.val(S, "W1045"))],
            "choki": num(b.val(S, "U1046")),
            "kashi": [num(b.val(S, f"{c}1048")) for c in ("U", "W", "Y", "AA", "AB")],
            "pgRoyalty": num(b.val(S, "S1050")),
        },
    }


# ---------------------------------------------------------------- ①大工
def extract_daiku(b):
    S = "①大工"
    out = {}
    for r in list(range(4, 27)) + list(range(30, 38)) + list(range(42, 57)):
        name = txt(b.val(S, f"B{r}")) or txt(b.val(S, f"C{r}"))
        if not name:
            continue
        F = num(b.val(S, f"F{r}"))
        G = b.fml(S, f"G{r}")
        if F is None and num(G) is not None:  # 仮床・気密テープは金額直書き
            F = num(G)
        out[r] = {"name": txt(b.val(S, f"B{r}")), "spec": txt(b.val(S, f"C{r}")),
                  "unit": txt(b.val(S, f"E{r}")), "price": F}
    return out


# ---------------------------------------------------------------- ②ｻｯｼ
def extract_sash(b):
    S = "②ｻｯｼ"
    items = []
    cat = "標準"
    for r in range(4, 368):
        A, B, C = txt(b.val(S, f"A{r}")), txt(b.val(S, f"B{r}")), txt(b.val(S, f"C{r}"))
        D = b.val(S, f"D{r}")
        if A and not C and not num(D):
            cat = A
            continue
        if not C and not B:
            continue
        kf = b.fml(S, f"K{r}") or ""
        extra = 156 if re.search(r"\+\s*156", str(kf)) else 0
        std = {}
        for col, glass in (("V", "clear"), ("W", "frost"), ("X", "rac")):
            f = b.fml(S, f"{col}{r}")
            if not is_formula(f):
                continue
            m = re.search(r'="PG",\s*0\s*,\s*(\d+)\)', f)
            if m:
                std.setdefault("normal", {})[glass] = int(m.group(1))
            m = re.search(r'="PG",\s*(\d+)\s*,\s*""\)', f)
            if m:
                std.setdefault("PG", {})[glass] = int(m.group(1))
        label = A if (A and C) else ""
        items.append({
            "row": r, "cat": cat, "label": label, "type": B, "size": C,
            "list": num(D), "rate": num(b.val(S, f"E{r}")),
            "listF": num(b.val(S, f"F{r}")), "rateF": num(b.val(S, f"G{r}")),
            "screen": num(b.val(S, f"H{r}")), "rateS": num(b.val(S, f"I{r}")),
            "other": num(b.val(S, f"J{r}")), "extra": extra,
            "madodai": num(b.val(S, f"M{r}")), "daiku": num(b.val(S, f"N{r}")),
            "angle": num(b.val(S, f"O{r}")), "order": num(b.val(S, f"R{r}")),
            "std": std or None,
            "note": "製作不可" if "製作不可" in C else "",
        })
    return items


# ---------------------------------------------------------------- ③既建
def extract_doors(b):
    S = "③既建"
    items = []
    group, frame, dtype = "", "", ""
    for r in range(5, 235):
        A, B, C, D = (txt(b.val(S, f"{c}{r}")) for c in "ABCD")
        if A:
            group = A
        if C:
            dtype = C
        K = num(b.val(S, f"K{r}"))
        if B == "■個別契約時見積":
            group = "個別契約時見積"
            continue
        if K is None and num(b.val(S, f"L{r}")) is None:
            continue
        std = {}
        vf = b.fml(S, f"V{r}")
        if is_formula(vf):
            if re.search(r'="PG",\s*0\s*,\s*1\)', vf):
                std["normal"] = 1
        # PG標準は契約内訳の P列数量（V14..V17 は式が無いので契約内訳から）
        items.append({
            "row": r, "group": group, "code": B if B.startswith("WD") else "",
            "frame": B if not B.startswith("WD") else D, "type": dtype,
            "w": num(b.val(S, f"F{r}")), "h": num(b.val(S, f"J{r}")),
            "list": K, "leaves": num(b.val(S, f"P{r}")) or 1,
            "L": num(b.val(S, f"L{r}")), "frameP": num(b.val(S, f"M{r}")),
            "handle": num(b.val(S, f"N{r}")), "closer": num(b.val(S, f"O{r}")),
            "daiku": num(b.val(S, f"R{r}")) or 0,
            "cost": num(b.val(S, f"Q{r}")),  # 原価（加工無）
            "stdCost": num(b.val(S, f"S{r}")),  # 標準仕様（デザイン・錠込み）原価
            "design": txt(b.val(S, f"W{r}")), "lock": txt(b.val(S, f"X{r}")),
            "std": std or None,
        })
    # PG標準数量（契約内訳 379-382 P列）
    pg_rows = {379: 14, 380: 15, 381: 16, 382: 17}
    for kr, dr in pg_rows.items():
        q = num(b.val("契約内訳", f"P{kr}"))
        for it in items:
            if it["row"] == dr and q:
                it["std"] = dict(it["std"] or {}, PG=q)
    designs = []
    for r in range(244, 311):
        j, k = txt(b.val(S, f"J{r}")), num(b.val(S, f"K{r}"))
        if j and k is not None:
            designs.append({"code": j, "list": k})
    locks = []
    for r in (311, 312):
        locks.append({"code": txt(b.val(S, f"J{r}")), "name": txt(b.val(S, f"I{r}")),
                      "list": num(b.val(S, f"K{r}"))})
    install = {"cost": num(b.val(S, "L53")) or 25000}
    net = {"base": num(b.val(S, "Q237")), "handleLess": num(b.val(S, "Q238")),
           "cut": num(b.val(S, "Q239")), "over": num(b.val(S, "Q240"))}
    return items, designs, locks, install, net


# ---------------------------------------------------------------- ④造建
def extract_zosaku(b):
    S = "④造建"
    items, install = [], []
    dtype = ""
    for r in range(4, 93):
        C = txt(b.val(S, f"C{r}"))
        if C:
            dtype = C
        H = num(b.val(S, f"H{r}"))
        if H is None:
            continue
        rec = {"row": r, "type": dtype, "sym": txt(b.val(S, f"D{r}")),
               "w": txt(b.val(S, f"E{r}")), "h": txt(b.val(S, f"F{r}")),
               "spec": txt(b.val(S, f"G{r}")), "cost": H,
               "hw": num(b.val(S, f"I{r}")) or 0, "daiku": num(b.val(S, f"J{r}")) or 0}
        if dtype == "搬入取付費":
            install.append({"row": r, "range": rec["spec"], "cost": H})
        else:
            items.append(rec)
    return items, install


# ---------------------------------------------------------------- ⑤固棚
def extract_shelves(b):
    S = "⑤固棚"
    items = []
    kind = ""
    for r in range(4, 38):
        A = txt(b.val(S, f"A{r}"))
        if A:
            kind = A
        B = txt(b.val(S, f"B{r}"))
        if not B:
            continue
        items.append({"row": r, "kind": kind, "size": B,
                      "board": num(b.val(S, f"C{r}")) or 0, "side": num(b.val(S, f"D{r}")) or 0,
                      "pipe": num(b.val(S, f"E{r}")) or 0, "hanger": num(b.val(S, f"F{r}")) or 0,
                      "install": num(b.val(S, f"I{r}")) or 0,
                      "std": num(b.val(S, f"O{r}")) or 0})
    return items


# ---------------------------------------------------------------- ⑥可棚
def extract_kadou(b):
    S = "⑥可棚"
    items = []
    label = ""
    rate = num(b.val(S, "F2"))
    for r in range(6, 65):
        A = txt(b.val(S, f"A{r}"))
        if A:
            label = A
        D = num(b.val(S, f"D{r}"))
        if D is None:
            continue
        F = num(b.val(S, f"F{r}")) or 0
        G = num(b.val(S, f"G{r}")) or 0
        H = num(b.val(S, f"H{r}")) or 0
        items.append({"row": r, "label": label, "b": txt(b.val(S, f"B{r}")),
                      "len": num(b.val(S, f"C{r}")), "cost": F + G, "side": H,
                      "kind": "board" if r <= 42 else "hw"})
    return {"items": items, "boardRate": rate, "install": num(b.val("①大工", "F17")) or 3000}


# ---------------------------------------------------------------- ⑦電気 / ⑧配管
def col_letter(i):
    return openpyxl.utils.get_column_letter(i)


def extract_electric(b):
    S = "⑦電気"
    items = []
    for c in range(3, 27):
        L = col_letter(c)
        name = txt(b.val(S, f"{L}4"))
        if not name:
            continue
        items.append({"col": L, "name": name, "sub": txt(b.val(S, f"{L}5")),
                      "std": num(b.val(S, f"{L}35")) or 0, "cost": num(b.val(S, f"{L}37")) or 0})
    dai1 = []
    for r in range(42, 46):
        dai1.append({"key": txt(b.val(S, f"AF{r}")), "floors": txt(b.val(S, f"AG{r}")),
                     "labor": num(b.val(S, f"AH{r}")), "unit": num(b.val(S, f"AI{r}")),
                     "maker": txt(b.val(S, f"AJ{r}"))})
    return {"items": items, "dai1": dai1}


def extract_plumbing(b):
    S = "⑧配管"
    items = []
    for c in range(3, 11):
        L = col_letter(c)
        name = txt(b.val(S, f"{L}4"))
        if not name:
            continue
        items.append({"col": L, "name": name, "std": num(b.val(S, f"{L}31")) or 0,
                      "cost": num(b.val(S, f"{L}33")) or 0})
    return items


# ---------------------------------------------------------------- 入力用
def extract_input(b):
    S = "入力用"
    cities = []
    for r in range(2, 6):
        n = txt(b.val(S, f"L{r}"))
        if n:
            cities.append({"name": n, "chokiFee": num(b.val(S, f"M{r}"))})
    # 原価率：入力用!B90 の式 =IF(D90<>"",D90,IF(B9="簡易",率,IF(B9="紹介",率,...,既定の率))) から区分ごとの率を読む
    f = str(b.fml(S, "B90") or "")
    rates = {k: int(v) for k, v in re.findall(r'B9="([^"]+)",(\d+)', f)}
    m = re.search(r',(\d+)\)+\s*$', f)
    rates["通常"] = int(m.group(1)) if m else None
    return {"cities": cities, "kakuninCost": num(b.val(S, "B5")),
            "kubunRate": rates}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsm")
    here = os.path.dirname(os.path.abspath(__file__))
    ap.add_argument("-o", "--out", default=os.path.join(here, "..", "js", "master-data.js"))
    a = ap.parse_args()

    b = Book(a.xlsm)
    doors, designs, locks, door_install, door_net = extract_doors(b)
    zosaku, zosaku_install = extract_zosaku(b)
    data = {
        "meta": {"source": os.path.basename(a.xlsm),
                 "extractedAt": datetime.datetime.now().strftime("%Y-%m-%d %H:%M")},
        "input": extract_input(b),
        "kouji": extract_kouji(b),
        "koujiExtra": extract_kouji_extra(b),
        "daiku": extract_daiku(b),
        "sash": extract_sash(b),
        "doors": doors, "doorDesigns": designs, "doorLocks": locks,
        "doorInstall": door_install, "doorNet": door_net,
        "zosaku": zosaku, "zosakuInstall": zosaku_install,
        "shelves": extract_shelves(b),
        "kadou": extract_kadou(b),
        "electric": extract_electric(b),
        "plumbing": extract_plumbing(b),
    }
    js = ("/* 自動生成ファイル: tools/extract_master.py で Excel から抽出。手で編集しないでください。 */\n"
          "var MASTER_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
          "if (typeof module !== 'undefined') module.exports = MASTER_DATA;\n")
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out, "w", encoding="utf-8") as fp:
        fp.write(js)
    print(f"wrote {a.out}: 工事明細 {len(data['kouji'])}行 / サッシ {len(data['sash'])} / "
          f"既製建具 {len(doors)} / 造作建具 {len(zosaku)} / 固定棚 {len(data['shelves'])} / "
          f"可動棚 {len(data['kadou']['items'])} / 電気 {len(data['electric']['items'])} / "
          f"配管 {len(data['plumbing'])}")


if __name__ == "__main__":
    main()
